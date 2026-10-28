"""Build synthetic import documents with the bundled document runtime; never edits references."""

from __future__ import annotations
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "演示导入文件" / "综合演示"
PROTOCOL = [
    ("heading", "项目与输入条件"),
    ("body", "1.1 本项目面向6×4重型牵引车，配套设备为商用车辆传动总成。"),
    ("body", "2.1 整车总质量49 t。"),
    ("body", "2.2 发动机最大扭矩2600 N·m。"),
    ("body", "2.3 发动机最大功率397 kW / 1900 r/min。"),
    ("heading", "性能与安装要求"),
    ("body", "3.1 采用12档AMT结构。"),
    ("body", "3.2 一档速比不低于14.5。"),
    ("body", "3.3 额定输入扭矩不低于发动机最大扭矩。"),
    ("body", "3.4 最高档为直接档或超速档，具体形式由双方进一步明确。"),
    ("body", "3.5 需配置取力器接口。"),
    ("body", "4.1 总成重量不超过300 kg。"),
    ("body", "4.2 离合器壳接口按主机厂合成图纸DEMO-DRW-001执行。"),
    ("break", ""),
    ("heading", "工况与可靠性要求"),
    ("body", "5.1 工作温度为-40～50 ℃。"),
    ("body", "5.2 配套使用条件为山区重载运输，按登记工况制定验证计划。"),
    ("body", "6.1 台架耐久按双方商定标准执行，详细判据待补充确认。"),
    ("body", "6.2 整机噪声不超过80 dB(A)。"),
    ("heading", "维护与适用依据"),
    ("body", "7.1 润滑油采用75W-90合成齿轮油。"),
    ("body", "7.2 换油周期不低于10万km。"),
    (
        "body",
        "8.1 适用标准采用合成评审规范DEMO-STD-001，正式工程依据由项目双方另行核准。",
    ),
    ("heading", "试验与交付验收"),
    ("body", "9.1 验收方法为台架验证与文件评审，提交带版本引用的记录后由工程师审核。"),
    (
        "body",
        "9.2 本文件只规定输入条件。实际执行记录、样件实测数据和协作周报应分别登记并保留证据。",
    ),
    (
        "body",
        "9.3 本文件及配套资料均为合成演示材料，用于本地软件流程验证，不作为真实产品合同、制造规格或认证报告。",
    ),
]


def write_json(name: str, data: object) -> None:
    (OUT / name).write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "01_综合技术协议_V1.docx"
    if path.exists():
        raise SystemExit("目标协议已存在，请人工核对后另选版本，生成器不覆盖原文件")
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    for name, size in [
        ("Normal", 11.5),
        ("Title", 21),
        ("Heading 1", 14),
        ("Subtitle", 10),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_after = Pt(9)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Subtitle"].font.italic = False
    doc.add_paragraph("重型变速器综合技术协议", style="Title")
    doc.add_paragraph(
        "编号 DEMO TX 001   版本 V1   日期 2026年10月23日", style="Subtitle"
    )
    doc.add_paragraph(
        "本协议用于智研引擎本地研制流程演示。供需双方以本文列出的条款作为输入，未定事项通过配套澄清单补充后再审核。全部参数和对象为合成数据。"
    )
    for kind, text in PROTOCOL:
        if kind == "break":
            doc.add_page_break()
        elif kind == "heading":
            doc.add_paragraph(text, style="Heading 1")
        else:
            doc.add_paragraph(text)
    doc.core_properties.author = "智研引擎项目组"
    doc.core_properties.title = "重型变速器综合技术协议"
    doc.core_properties.comments = "本地合成演示输入"
    doc.save(path)
    write_json(
        "协议正文与编号映射.json",
        {
            "data_kind": "synthetic",
            "protocol_id": "DEMO-TX-001",
            "version": "V1",
            "paragraphs": PROTOCOL,
            "display_code_map": {
                "engineMaxTorque": "REQ-P01-B-003",
                "gearboxRatedInputTorque": "REQ-P01-C-002",
                "gearCount": "REQ-P01-C-001",
                "transmissionType": "REQ-P01-C-001-T",
                "pto": "REQ-P01-C-006",
                "assemblyWeight": "REQ-P01-D-002",
                "durability": "REQ-P01-F-001",
            },
        },
    )
    (OUT / "02_澄清单与变更说明.txt").write_text(
        """全部内容均为合成演示记录。
CN-01：最高档形式明确为超速档。修订最高档需求value_text为“超速档”，source_excerpt填本澄清记录；关闭原模糊问题并重新确认需求。
CN-02：耐久依据为合成补充规程DEMO-DUR-001，等效台架循环1000 h无失效。先登记澄清证据，补齐并确认用例判据；F-001仍计入固定5条验证范围。
BG-01：发动机最大扭矩由2600改为2800 N·m。只修订源参数B-003，系统更新独立派生要求C-002；基型JX-12A额定2600不改变。旧试验仍保留，新用例须重新基线和执行。
合成试验边界：静扭实际加载2800且后备系数2.4不通过；加载2600即使系数2.6也不算新版通过；加载2800且系数2.6符合静扭本项。
温度补充说明：20℃处于-40～50℃区间只证明该点；需要边界覆盖时单独登记条件与证据，不以单点推断全范围。
NVH演示可另登记82 dB(A)的实测证据，必须呈现与80上限的偏差；本次固定验证范围仍为5条，不把额外说明混入已通过结论。
""",
        encoding="utf-8",
    )
    write_json(
        "03_零件选型样例.json",
        {
            "data_kind": "synthetic",
            "catalog": "数据库part_candidates，017迁移预置24型号",
            "engine_input": {
                "max_torque_nm": 2600,
                "power_kw": 397,
                "power_speed_rpm": 1900,
            },
            "base": "JX-12A",
            "selections": [
                {
                    "part_name": "主箱壳体",
                    "model_code": "DEMO-P01-3",
                    "reason": "3000 N·m / 16档合成匹配包络；材料与安装接口仍需工程师核对",
                },
                {
                    "part_name": "齿轮组",
                    "model_code": "DEMO-P02-3",
                    "reason": "加强包络示例",
                },
            ],
            "incompatible_example": {
                "model_code": "DEMO-P01-1",
                "reason": "2400 N·m / 8档不满足本协议输入，不允许选用",
            },
            "image_note": "界面使用本地工程示意图，产品照片为参考图，不是候选实物认证照片",
        },
    )
    write_json(
        "04_试验与评审样例.json",
        {
            "data_kind": "synthetic",
            "mode": "manual",
            "review_comment": "已依据协议及合成澄清单核对当前版本",
            "fmea_scores": {
                "FMEA-01": {"severity": 8, "occurrence": 4, "detection": 3},
                "FMEA-02": {"severity": 7, "occurrence": 3, "detection": 4},
                "FMEA-03": {"severity": 6, "occurrence": 3, "detection": 4},
                "default": {"severity": 6, "occurrence": 3, "detection": 4},
            },
            "fmea_basis": "工程师按演示量表人工指定；不按行号编造分数，不代表实测风险频度",
            "durability_criteria": "合成补充规程DEMO-DUR-001：等效台架循环1000 h无失效",
            "manual_run": {
                "conditions": "合成台架/静态检查工况",
                "manual_result": "pass",
                "evidence": "合成执行记录DEMO-TR-01，参数与用例逐项核对",
            },
            "static_torsion": {
                "initial_pass": {"load_basis_numeric": 2600, "backup_coefficient": 2.6},
                "wrong_load": {"load_basis_numeric": 2600, "backup_coefficient": 2.6},
                "new_fail": {"load_basis_numeric": 2800, "backup_coefficient": 2.4},
                "new_pass": {"load_basis_numeric": 2800, "backup_coefficient": 2.6},
            },
        },
    )
    delivery = {
        "display_code": "SD-DEMO-001",
        "supplier": "合成样件提供方",
        "part_display_code": "P-03",
        "drawing_display_code": "DEMO-DRW-001",
        "drawing_revision": "A",
        "delivered_at": "2026-10-23",
        "quantity": 1,
        "notes": "合成尺寸基准，不由系统读取CAD自动识别",
        "drawing_parameters": [
            {
                "parameter_key": "holePosition",
                "parameter_name": "孔位",
                "expected_value": 120,
                "tolerance": 0.5,
                "unit": "mm",
                "source_note": "配套示意图尺寸1",
            },
            {
                "parameter_key": "shaftDiameter",
                "parameter_name": "轴径",
                "expected_value": 45,
                "tolerance": 0.02,
                "unit": "mm",
                "source_note": "配套示意图尺寸2",
            },
        ],
    }
    comparison = {
        "parameters": [
            {"parameter_key": "holePosition", "measured_value": 120.2, "unit": "mm"},
            {"parameter_key": "shaftDiameter", "measured_value": 45.01, "unit": "mm"},
        ],
        "visual_check": {
            "mode": "manual",
            "result": "pass",
            "note": "合成目视记录，由操作者明确录入",
        },
        "evidence": "合成尺寸记录DEMO-MEASURE-001",
    }
    fail = {
        **comparison,
        "parameters": [
            {"parameter_key": "holePosition", "measured_value": 121, "unit": "mm"},
            {"parameter_key": "shaftDiameter", "measured_value": 45.01, "unit": "mm"},
        ],
        "evidence": "合成超差记录DEMO-MEASURE-002",
    }
    write_json(
        "05_样件与实测样例.json",
        {
            "data_kind": "synthetic",
            "mode": "manual",
            "delivery": delivery,
            "file": "07_样件尺寸示意.svg",
            "comparison_pass": comparison,
            "comparison_fail": fail,
        },
    )
    write_json(
        "06_协同记录与进度样例.json",
        {
            "data_kind": "synthetic",
            "topic": {
                "name": "合成设计冻结交付协同",
                "expected_participants": ["A 主机方", "B 研发方"],
                "window_start": "2026-10-20T00:00:00+08:00",
                "window_end": "2026-10-24T00:00:00+08:00",
            },
            "records": [
                {
                    "external_id": "DEMO-MSG-01",
                    "record_type": "message",
                    "sender": "A 主机方",
                    "content_summary": "请确认交付资料",
                    "sent_at": "2026-10-20T09:00:00+08:00",
                    "reply_to_external_id": None,
                },
                {
                    "external_id": "DEMO-MSG-02",
                    "record_type": "message",
                    "sender": "B 研发方",
                    "content_summary": "回复资料核对状态",
                    "sent_at": "2026-10-20T11:00:00+08:00",
                    "reply_to_external_id": "DEMO-MSG-01",
                },
            ],
            "progress": [
                {
                    "participant": "A 主机方",
                    "milestone": "设计冻结",
                    "planned_percent": 80,
                    "actual_percent": 90,
                    "artifact_version": "V2",
                    "reported_at": "2026-10-23T16:00:00+08:00",
                    "evidence": "合成周报A",
                },
                {
                    "participant": "B 研发方",
                    "milestone": "设计冻结",
                    "planned_percent": 80,
                    "actual_percent": 30,
                    "artifact_version": "V1",
                    "reported_at": "2026-10-23T16:00:00+08:00",
                    "evidence": "合成周报B",
                },
            ],
            "risk_acceptance_note": "合成排练中明确接受已登记进度差和交付版本风险，后续按计划复核；不代表产品不合格结论",
        },
    )
    (OUT / "07_样件尺寸示意.svg").write_text(
        """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 900 500"><rect width="900" height="500" fill="#f4f8fc"/><g stroke="#2e5677" fill="none" stroke-width="3"><rect x="180" y="100" width="540" height="260" rx="18"/><circle cx="300" cy="230" r="55"/><circle cx="600" cy="230" r="55"/><path d="M300 160V400M600 160V400M300 390H600M645 175H780M645 285H780M765 175V285" stroke-dasharray="8 5"/></g><g font-family="Microsoft YaHei, sans-serif" fill="#284f70"><text x="40" y="48" font-size="24">合成样件尺寸基准 DEMO-DRW-001 版本A</text><text x="370" y="430" font-size="24">孔位 120 ± 0.5 mm</text><text x="650" y="245" font-size="18">轴径 45 ± 0.02</text><text x="40" y="480" font-size="18">示意图非比例制造图  数值以登记基准为准  无CAD自动识别</text></g></svg>""",
        encoding="utf-8",
    )
    print(path)


if __name__ == "__main__":
    main()
